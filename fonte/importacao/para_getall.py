# Converte a planilha do app (xlsx baixado do Drive) no formato do getAll (matrizes, touros, protocolos, eventos, config)
import openpyxl, json, datetime, sys
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
TIPOS = {'Nascimentos':'nascimento','Desmama':'desmama','Recria':'recria','IATF':'iatf','Diagnosticos':'diagnostico','Manejos':'manejo','Pesagens':'pesagem','Vendas':'venda','VendaItens':'vendaitem','Insumos':'insumo','Descartes':'descarte','Mortes':'morte'}
def val(v):
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d') if (v.hour, v.minute) == (0, 0) or True else v.isoformat()
    if isinstance(v, datetime.date): return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer(): return int(v)
    return '' if v is None else v
def tab(n, ncols=None):
    ws = wb[n]; cab = [c.value for c in ws[1]]
    if ncols: cab = cab[:ncols]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] in (None, ''): continue
        out.append({k: val(v) for k, v in zip(cab, r) if k})
    return out
mat = tab('Matrizes', 12)
ev = []
for aba, t in TIPOS.items():
    for e in tab(aba): e['tipo'] = t; ev.append(e)
cfg = {r['chave']: r['valor'] for r in tab('Config')}
json.dump({'versao': 13, 'matrizes': mat, 'touros': tab('Touros'), 'protocolos': tab('Protocolos'), 'eventos': ev, 'config': cfg}, open(sys.argv[2], 'w'), ensure_ascii=False, default=str)
print(len(mat), len(ev))
